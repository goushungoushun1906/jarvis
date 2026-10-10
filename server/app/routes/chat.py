"""JARVIS chat & conversation routes."""

from __future__ import annotations

import asyncio
import base64
import json
import logging
import os

from fastapi import APIRouter, File, HTTPException, Request, UploadFile
from fastapi.responses import JSONResponse, StreamingResponse

from .. import database as db
from .. import memory as mem
from .. import persona as persona_mod
from ..agent import agent_chat, agent_chat_stream, strip_tool_call_json
from ..approval import new_gate
from ..config import settings
from ..llm import LLMError, call_llm
from ..models import (
    ChatRequest,
    ConversationCreate,
    ConversationOut,
    ForkRequest,
    MessageOut,
)
from ..tools import registry as tool_registry

logger = logging.getLogger("jarvis")

router = APIRouter(prefix="/api", tags=["chat"])

# Phase 17.3 — SSE heartbeat: comments are emitted when no agent event has
# been produced for this long, keeping intermediaries (proxies, antivirus,
# the Electron layer) from killing an idle connection during long tool runs.
HEARTBEAT_INTERVAL_SECONDS = 15.0


def _sse(event: str, payload: dict) -> str:
    """Format one named SSE event.

    Every event carries BOTH an ``event:`` line (new contract, Phase 17.3)
    and a ``data:`` JSON line (legacy contract). The current frontend parses
    only ``data:`` lines, so old clients keep working unchanged; new clients
    can dispatch on ``event:``.
    """
    return f"event: {event}\ndata: {json.dumps(payload, ensure_ascii=False, default=str)}\n\n"


# ------------------------------------------------------------------
# Message format conversion helpers
# ------------------------------------------------------------------

def _message_to_llm_format(role: str, content: str | None,
                            images: list = None,
                            files: list = None) -> dict:
    """Convert app message format to OpenAI API format."""
    images = images or []
    files = files or []

    text_parts = []
    if content:
        text_parts.append(content)
    for f in files:
        f_dict = f if isinstance(f, dict) else f.model_dump()
        if f_dict.get("extracted_text"):
            text_parts.append(f"[文件 {f_dict['filename']} 的内容]:\n{f_dict['extracted_text']}")
    text = "\n\n".join(text_parts) if text_parts else ""

    if images:
        parts = []
        if text:
            parts.append({"type": "text", "text": text})
        for img in images:
            img_dict = img if isinstance(img, dict) else img.model_dump()
            parts.append({"type": "image_url", "image_url": {"url": img_dict["image_url"]["url"]}})
        return {"role": role, "content": parts}
    else:
        return {"role": role, "content": text}


def _db_message_to_llm_format(msg: dict) -> dict:
    """Convert database message to LLM format. Handles both plain text and JSON multimodal."""
    content_str = msg["content"]
    role = msg["role"]

    try:
        payload = json.loads(content_str)
        if isinstance(payload, dict) and ("images" in payload or "files" in payload):
            return _message_to_llm_format(
                role, payload.get("text"), payload.get("images"), payload.get("files")
            )
    except (json.JSONDecodeError, TypeError):
        pass

    return {"role": role, "content": content_str}


# ------------------------------------------------------------------
# Conversations CRUD
# ------------------------------------------------------------------

@router.get("/conversations", response_model=list[ConversationOut])
async def api_list_conversations(limit: int = 50, offset: int = 0):
    rows = await db.list_conversations(limit=limit, offset=offset)
    return rows


@router.post("/conversations", response_model=ConversationOut, status_code=201)
async def api_create_conversation(body: ConversationCreate):
    conv = await db.create_conversation(title=body.title)
    return {**conv, "message_count": 0}


@router.get("/conversations/{conversation_id}")
async def api_get_conversation(conversation_id: str):
    conv = await db.get_conversation(conversation_id)
    if conv is None:
        raise HTTPException(404, "Conversation not found")
    return conv


@router.delete("/conversations/{conversation_id}", status_code=204)
async def api_delete_conversation(conversation_id: str):
    deleted = await db.delete_conversation(conversation_id)
    if not deleted:
        raise HTTPException(404, "Conversation not found")


@router.patch("/conversations/{conversation_id}")
async def api_patch_conversation(conversation_id: str, request: Request):
    body = await request.json()
    title = body.get("title")
    if title is None:
        raise HTTPException(422, "Missing 'title' field")
    ok = await db.update_conversation_title(conversation_id, title)
    if not ok:
        raise HTTPException(404, "Conversation not found")
    conv = await db.get_conversation(conversation_id)
    return conv


@router.post("/conversations/{conversation_id}/fork", response_model=ConversationOut, status_code=201)
async def api_fork_conversation(conversation_id: str, body: ForkRequest):
    """Fork conversation from a specific message."""
    original = await db.get_conversation(conversation_id)
    if original is None:
        raise HTTPException(404, "Conversation not found")

    all_messages = await db.get_messages(conversation_id)
    fork_index = None
    for i, msg in enumerate(all_messages):
        if msg["id"] == body.message_id:
            fork_index = i
            break

    if fork_index is None:
        raise HTTPException(404, "Message not found in this conversation")

    messages_to_copy = all_messages[:fork_index + 1]
    fork_title = body.new_title or f"{original.get('title') or '对话'} (分支)"

    new_convo = await db.create_conversation(
        title=fork_title,
        parent_id=conversation_id,
        forked_from_message_id=body.message_id,
    )

    for msg in messages_to_copy:
        await db.add_message(new_convo["id"], msg["role"], msg["content"])

    return {**new_convo, "message_count": len(messages_to_copy)}


# ------------------------------------------------------------------
# Upload endpoints
# ------------------------------------------------------------------

@router.post("/upload/image")
async def upload_image(file: UploadFile = File(...)):
    """Upload image, return base64 data URL for preview and sending."""
    allowed_types = {"image/png", "image/jpeg", "image/gif", "image/webp"}
    if file.content_type not in allowed_types:
        raise HTTPException(400, f"不支持的图片类型: {file.content_type}")
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(400, "图片大小不能超过 10MB")
    b64 = base64.b64encode(content).decode("utf-8")
    data_url = f"data:{file.content_type};base64,{b64}"
    return {"data_url": data_url, "filename": file.filename, "mime_type": file.content_type}


@router.post("/upload/file")
async def upload_file(file: UploadFile = File(...)):
    """Upload file, extract text content. Supports TXT, PDF, MD, CSV, JSON, DOC, DOCX, XLS, XLSX."""
    allowed_extensions = {".txt", ".pdf", ".md", ".csv", ".json", ".doc", ".docx", ".xls", ".xlsx", ".pptx", ".ppt"}
    filename = file.filename or "unknown"
    ext = os.path.splitext(filename)[1].lower()
    logger.info("[upload_file] name=%s ext=%s content_type=%s", filename, ext, file.content_type)
    if ext not in allowed_extensions:
        raise HTTPException(400, f"不支持的文件类型: {ext}")
    content_bytes = await file.read()
    logger.info("[upload_file] read %d bytes", len(content_bytes))
    b64 = base64.b64encode(content_bytes).decode("utf-8")
    data_url = f"data:{file.content_type or 'application/octet-stream'};base64,{b64}"
    if len(content_bytes) > 20 * 1024 * 1024:
        raise HTTPException(400, "文件大小不能超过 20MB")

    extracted_text = ""
    try:
        if ext in {".txt", ".md", ".csv", ".json"}:
            extracted_text = content_bytes.decode("utf-8", errors="replace")
        elif ext == ".pdf":
            import fitz
            doc = fitz.open(stream=content_bytes, filetype="pdf")
            pages = []
            for page in doc:
                pages.append(page.get_text())
            doc.close()
            extracted_text = "\n\n".join(pages)
        elif ext in {".doc", ".docx"}:
            import docx
            import io
            try:
                doc = docx.Document(io.BytesIO(content_bytes))
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                tables_text = []
                for table in doc.tables:
                    rows = []
                    for row in table.rows:
                        rows.append(" | ".join(cell.text for cell in row.cells))
                    tables_text.append("\n".join(rows))
                extracted_text = "\n\n".join(paragraphs)
                if tables_text:
                    extracted_text += "\n\n[表格内容]:\n" + "\n\n".join(tables_text)
            except Exception:
                # Try antiword / textract for old .doc files
                if ext == ".doc":
                    try:
                        import subprocess, tempfile
                        with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as tf:
                            tf.write(content_bytes)
                            tmp_path = tf.name
                        result = subprocess.run(
                            ["antiword", tmp_path], capture_output=True, text=True, timeout=30
                        )
                        os.unlink(tmp_path)
                        if result.returncode == 0 and result.stdout.strip():
                            extracted_text = result.stdout.strip()
                        else:
                            raise RuntimeError("antiword failed")
                    except Exception:
                        raise RuntimeError(
                            "无法解析旧版 .doc 文件，请将其另存为 .docx 格式后再上传"
                        )
                else:
                    raise
            if not extracted_text:
                extracted_text = "[文档中未提取到可读文本]"
        elif ext in {".xls"}:
            # Old Excel format — use xlrd
            import xlrd
            import io
            wb = xlrd.open_workbook(file_contents=content_bytes)
            sheets_text = []
            for sheet in wb.sheets():
                rows = []
                for row_idx in range(sheet.nrows):
                    row_values = [str(sheet.cell_value(row_idx, col_idx)) for col_idx in range(sheet.ncols)]
                    row_text = " | ".join(row_values)
                    if row_text.strip():
                        rows.append(row_text)
                if rows:
                    sheets_text.append(f"[工作表: {sheet.name}]\n" + "\n".join(rows))
            extracted_text = "\n\n".join(sheets_text)
        elif ext in {".xlsx"}:
            # New Excel format — use openpyxl
            import openpyxl
            import io
            wb = openpyxl.load_workbook(io.BytesIO(content_bytes), read_only=True, data_only=True)
            sheets_text = []
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join(str(cell) if cell is not None else "" for cell in row)
                    if row_text.strip():
                        rows.append(row_text)
                if rows:
                    sheets_text.append(f"[工作表: {sheet_name}]\n" + "\n".join(rows))
            wb.close()
            extracted_text = "\n\n".join(sheets_text)
    except ImportError as e:
        logger.error("[upload_file] ImportError: %s", e)
        raise HTTPException(500, f"文件解析库未安装: {str(e)}")
    except Exception as e:
        logger.error("[upload_file] Parse error for %s: %s", filename, e)
        raise HTTPException(500, f"文件解析失败: {str(e)}")

    if len(extracted_text) > 50000:
        extracted_text = extracted_text[:50000] + f"\n\n[... 文本已截断，共 {len(extracted_text)} 字 ...]"

    return {"filename": filename, "mime_type": file.content_type or "", "extracted_text": extracted_text, "char_count": len(extracted_text), "data_url": data_url}


# ------------------------------------------------------------------
# Chat
# ------------------------------------------------------------------

async def _build_llm_messages(conversation_id: str | None, request_messages: list[dict]) -> list[dict]:
    """Prepend system prompt and optionally load conversation history."""
    # Keep system message last so we can use the user's latest message as the
    # memory retrieval query.
    messages: list[dict] = []

    if conversation_id:
        history = await db.get_messages(conversation_id, limit=settings.max_context_messages)
        for msg in history:
            messages.append(_db_message_to_llm_format(msg))

    messages.extend(request_messages)

    persona = await persona_mod.get_persona_config()
    system_prompt = await persona_mod.build_system_prompt(persona, messages=messages)
    messages.insert(0, {"role": "system", "content": system_prompt})
    return messages


async def _save_user_message(conversation_id: str, role: str, content: str | None,
                              images: list = None, files: list = None) -> None:
    """Save user message to database. Multimodal content serialized as JSON."""
    images = images or []
    files = files or []
    if images or files:
        payload = {"text": content or ""}
        if images:
            payload["images"] = [
                {"image_url": (img if isinstance(img, dict) else img.model_dump())["image_url"], "filename": (img if isinstance(img, dict) else img.model_dump()).get("filename", "")}
                for img in images
            ]
        if files:
            payload["files"] = [
                {"filename": (f if isinstance(f, dict) else f.model_dump())["filename"], "extracted_text": (f if isinstance(f, dict) else f.model_dump())["extracted_text"]}
                for f in files
            ]
        await db.add_message(conversation_id, role, json.dumps(payload, ensure_ascii=False))
    else:
        await db.add_message(conversation_id, role, content or "")


async def _save_assistant_message(conversation_id: str, content: str) -> None:
    await db.add_message(conversation_id, "assistant", content)


@router.post("/chat")
async def api_chat(req: ChatRequest):
    """Non-streaming chat completion with optional conversation persistence.
    Uses the agent loop for tool support."""
    request_messages = []
    for m in req.messages:
        msg_dict = _message_to_llm_format(m.role, m.content, m.images, m.files)
        request_messages.append(msg_dict)
    conversation_id = req.conversation_id

    try:
        if conversation_id:
            # Verify conversation exists
            conv = await db.get_conversation(conversation_id)
            if conv is None:
                raise HTTPException(404, "Conversation not found")

        messages = await _build_llm_messages(conversation_id, request_messages)

        # Use agent loop (with tool support) for non-streaming.
        # A6 — this endpoint cannot render an approval dialog (the response is
        # not sent until the loop ends), so gated tools are refused outright
        # rather than hanging until the approval timeout expires.
        content, tool_log = await agent_chat(messages, new_gate(auto_deny=True))

        # Persist if conversation_id provided
        if conversation_id:
            # Save each user message from the request
            for m in req.messages:
                if m.role == "user":
                    await _save_user_message(conversation_id, m.role, m.content, m.images, m.files)
            await _save_assistant_message(conversation_id, content)

            # Auto-extract memories from user messages
            for m in req.messages:
                if m.role == "user":
                    try:
                        text_for_memory = m.content or ""
                        if m.files:
                            text_for_memory += "\n" + " ".join(
                                f.get("extracted_text", "") for f in m.files
                            )
                        await mem.extract_memories_from_message(text_for_memory, conversation_id)
                    except Exception as e:
                        logger.error("Memory extraction error: %s", e)

        return {
            "content": content,
            "role": "assistant",
            "conversation_id": conversation_id,
            "tool_calls": tool_log if tool_log else None,
        }

    except LLMError as e:
        return JSONResponse(
            status_code=e.status_code,
            content={"error": e.message, "detail": e.detail},
        )


@router.post("/chat/stream")
async def api_chat_stream(req: ChatRequest):
    """Streaming chat completion via SSE with optional conversation persistence."""
    request_messages = []
    for m in req.messages:
        msg_dict = _message_to_llm_format(m.role, m.content, m.images, m.files)
        request_messages.append(msg_dict)
    conversation_id = req.conversation_id

    try:
        if conversation_id:
            conv = await db.get_conversation(conversation_id)
            if conv is None:
                raise HTTPException(404, "Conversation not found")

        messages = await _build_llm_messages(conversation_id, request_messages)

        # Save user messages before streaming
        if conversation_id:
            for m in req.messages:
                if m.role == "user":
                    await _save_user_message(conversation_id, m.role, m.content, m.images, m.files)

    except LLMError as e:
        return JSONResponse(
            status_code=e.status_code,
            content={"error": e.message, "detail": e.detail},
        )

    async def event_generator():
        """Phase 17.3 — consume AgentEvents and emit named SSE frames.

        Mapping (AgentEvent → SSE):
            text       → event: text        data: {"content": delta}
            tool_call  → event: tool_call   data: {"tool_call": {...}}
            tool_result→ event: tool_result data: {"tool_result": {...}}
            done       → event: done        data: {"conversation_id", "tool_calls"}
            error      → event: error       data: {"error": msg}

        ``data: [DONE]`` terminates the stream in both success and error paths.
        The assistant message is persisted only on the ``done`` path.
        """
        full_content: list[str] = []
        tool_log: list[dict] = []

        def _record_tool_log_entry(call_id: str, name: str, args: dict) -> dict:
            entry = {
                "call_id": call_id,
                "tool": name,
                "args": args,
                "result": "",
                "success": False,
            }
            tool_log.append(entry)
            return entry

        # A6 — one gate per stream. Every tool marked requires_approval blocks
        # here until POST /api/tools/approve resolves it (or it times out).
        gate = new_gate()

        agen = agent_chat_stream(messages, approval_gate=gate)
        next_task: asyncio.Task | None = None
        try:
            while True:
                # Drive the agent loop as a *task* and wait on it with a
                # timeout instead of asyncio.wait_for(agen.__anext__(), ...).
                #
                # wait_for cancels the inner coroutine when the deadline hits,
                # and cancelling __anext__ throws CancelledError *inside* the
                # generator frame — which closes the whole agent loop. That
                # was survivable while every round finished in a few seconds,
                # but A6 lets a round legitimately block for minutes waiting on
                # a human, and one heartbeat would have killed the stream.
                # asyncio.wait() leaves the task running across timeouts.
                if next_task is None:
                    next_task = asyncio.ensure_future(agen.__anext__())
                done, _pending = await asyncio.wait(
                    {next_task}, timeout=HEARTBEAT_INTERVAL_SECONDS
                )
                if not done:
                    yield ": ping\n\n"
                    continue
                try:
                    ev = next_task.result()
                except StopAsyncIteration:
                    break
                finally:
                    next_task = None

                if ev.kind == "text":
                    full_content.append(ev.delta)
                    yield _sse("text", {"content": ev.delta})

                elif ev.kind == "tool_call":
                    tool = tool_registry.get(ev.name)
                    # Prefer the loop's own judgement (it knows whether a gate
                    # is attached); fall back to class metadata for streams
                    # that run without a gate.
                    requires_approval = bool(
                        ev.requires_approval
                        if ev.requires_approval is not None
                        else getattr(tool, "requires_approval", False)
                    )
                    _record_tool_log_entry(ev.call_id, ev.name, ev.args)
                    yield _sse(
                        "tool_call",
                        {
                            "tool_call": {
                                "id": ev.call_id,
                                "tool": ev.name,
                                "args": ev.args,
                                "status": (
                                    "awaiting_approval"
                                    if requires_approval
                                    else "running"
                                ),
                                "requires_approval": requires_approval,
                            }
                        },
                    )

                elif ev.kind == "tool_result":
                    status = "success" if ev.success else "error"
                    for entry in reversed(tool_log):
                        if entry["call_id"] == ev.call_id and entry["result"] == "":
                            entry["result"] = ev.result
                            entry["success"] = ev.success
                            break
                    else:
                        _record_tool_log_entry(ev.call_id, ev.name, ev.args)
                        tool_log[-1]["result"] = ev.result
                        tool_log[-1]["success"] = ev.success
                    yield _sse(
                        "tool_result",
                        {
                            "tool_result": {
                                "id": ev.call_id,
                                "tool": ev.name,
                                "result": ev.result,
                                "success": ev.success,
                                "status": status,
                                "duration_ms": ev.duration_ms,
                                "metadata": ev.metadata,
                            }
                        },
                    )

                elif ev.kind == "done":
                    yield _sse(
                        "done",
                        {
                            "conversation_id": conversation_id,
                            "tool_calls": tool_log or None,
                        },
                    )
                    break

                elif ev.kind == "error":
                    logger.error("Agent stream error: %s", ev.message)
                    yield _sse("error", {"error": ev.message})
                    yield "data: [DONE]\n\n"
                    return

            yield "data: [DONE]\n\n"
        except asyncio.CancelledError:
            # Client disconnected mid-stream; persist nothing partial.
            # aclose() (in the finally below) is what actually stops the agent
            # loop — without it the LLM request and any in-flight tool keep
            # running detached until GC happens to collect the generator.
            logger.info("Client disconnected mid-stream; tearing down agent loop.")
            raise
        except Exception as e:
            logger.exception("Stream unexpected error: %s", e)
            yield _sse("error", {"error": str(e)})
            yield "data: [DONE]\n\n"
            return
        finally:
            if next_task is not None and not next_task.done():
                next_task.cancel()
            gate.close()
            try:
                await agen.aclose()
            except (RuntimeError, asyncio.CancelledError):
                pass
            except Exception:  # pragma: no cover - defensive
                logger.debug("agent stream close error", exc_info=True)

        # Save the full assistant response after streaming completes
        if conversation_id and full_content:
            # Strip tool-call JSON the model embedded in the text body, so the
            # persisted history matches what the user saw.
            await _save_assistant_message(
                conversation_id, strip_tool_call_json("".join(full_content))
            )

        # Auto-extract memories from user messages after streaming
        if conversation_id:
            for m in req.messages:
                if m.role == "user":
                    try:
                        text_for_memory = m.content or ""
                        if m.files:
                            text_for_memory += "\n" + " ".join(
                                f.get("extracted_text", "") for f in m.files
                            )
                        await mem.extract_memories_from_message(text_for_memory, conversation_id)
                    except Exception as e:
                        logger.error("Memory extraction error: %s", e)

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            # Disable proxy buffering so events reach the browser immediately
            "X-Accel-Buffering": "no",
        },
    )
